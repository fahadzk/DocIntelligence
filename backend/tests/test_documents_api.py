from io import BytesIO
from pathlib import Path
from uuid import UUID
import sqlite3

import fitz
from docx import Document as WordDocument

from app.config.settings import get_settings
from app.main import get_document_service
from app.infrastructure.migrations import migrate


def project(client, name="Research"):
    return client.post("/api/projects", json={"name": name}).json()["id"]


def upload(client, project_id, filename, contents):
    response = client.post(f"/api/projects/{project_id}/documents", files=[("files", (filename, contents, "application/octet-stream"))])
    assert response.status_code == 200, response.text
    return response.json()[0]


def pdf_bytes(texts):
    pdf = fitz.open()
    for text in texts:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    contents = pdf.tobytes()
    pdf.close()
    return contents


def docx_bytes():
    document = WordDocument()
    document.add_heading("Overview", level=1)
    document.add_paragraph("A useful paragraph.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_supported_formats_and_source_locations(client):
    project_id = project(client)
    examples = [
        ("report.pdf", pdf_bytes(["First page", "Second page"]), "First page"),
        ("report.docx", docx_bytes(), "A useful paragraph."),
        ("notes.txt", b"First paragraph." + bytes([10, 10]) + b"Second paragraph.", "Second paragraph."),
        ("notes.md", b"# Heading" + bytes([10, 10]) + b"Markdown body.", "Markdown body."),
    ]
    for filename, contents, expected in examples:
        imported = upload(client, project_id, filename, contents)
        document = client.get(f"/api/projects/{project_id}/documents/{imported['document']['id']}").json()
        assert document["status"] == "ready"
        content = client.get(f"/api/projects/{project_id}/documents/{document['id']}/content").json()
        assert expected in " ".join(segment["text"] for segment in content["segments"])
        if filename.endswith(".pdf"):
            assert document["page_count"] == 2
            assert [segment["page_number"] for segment in content["segments"]] == [1, 2]
        if filename.endswith(".docx"):
            assert content["segments"][0]["kind"] == "heading"
            assert content["segments"][1]["label"] == "Overview"
        if filename.endswith(".md"):
            assert content["segments"][0]["kind"] == "heading"


def test_failures_are_distinct_and_persisted(client):
    project_id = project(client)
    cases = [
        ("scan.pdf", pdf_bytes([""]), "NO_EXTRACTABLE_TEXT"),
        ("locked.pdf", password_pdf(), "PASSWORD_REQUIRED"),
        ("broken.pdf", b"%PDF-broken", "CORRUPT_FILE"),
        ("empty.txt", b"", "EMPTY_FILE"),
        ("unknown.exe", b"x", "UNSUPPORTED_TYPE"),
        ("bad.txt", bytes([255, 254]), "UNREADABLE_FILE"),
    ]
    for filename, contents, code in cases:
        outcome = upload(client, project_id, filename, contents)
        if outcome["document"]:
            document_id = outcome["document"]["id"]
            item = client.get(f"/api/projects/{project_id}/documents/{document_id}").json()
            assert item["status"] == "failed"
            assert item["error_code"] == code
            assert item["error_message"]
        else:
            assert outcome["error_code"] == code


def password_pdf():
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Secret")
    contents = pdf.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    pdf.close()
    return contents


def test_duplicate_collision_isolation_and_deletion(client, tmp_path: Path):
    first = project(client, "First")
    second = project(client, "Second")
    source = tmp_path / "same.txt"
    source.write_text("Version one", encoding="utf-8")
    one = upload(client, first, source.name, source.read_bytes())["document"]
    duplicate = upload(client, first, "copy.txt", source.read_bytes())
    assert duplicate["error_code"] == "DUPLICATE_DOCUMENT"
    two = upload(client, first, source.name, b"Version two")["document"]
    assert one["id"] != two["id"]
    assert upload(client, second, source.name, source.read_bytes())["document"]
    assert client.get(f"/api/projects/{second}/documents/{one['id']}").status_code == 404
    assert client.get(f"/api/projects/{second}/documents/{one['id']}/original").status_code == 404
    original = get_settings().data_dir / "documents" / first / one["id"] / "original.txt"
    assert original.is_file()
    assert client.delete(f"/api/projects/{first}/documents/{one['id']}").status_code == 204
    assert not original.exists()
    assert source.is_file()
    assert len(client.get(f"/api/projects/{first}/documents").json()) == 1
    assert len(client.get(f"/api/projects/{second}/documents").json()) == 1


def test_restart_persistence_interruption_and_integrity(client):
    project_id = project(client)
    document = upload(client, project_id, "persist.txt", b"Persistent text")["document"]
    document_id = document["id"]
    from app.main import get_project_service
    get_project_service.cache_clear()
    get_document_service.cache_clear()
    assert client.get(f"/api/projects/{project_id}/documents/{document_id}/content").json()["segments"][0]["text"] == "Persistent text"
    repository = get_document_service().repository
    repository.update(UUID(project_id), UUID(document_id), "processing", "extracting")
    get_document_service.cache_clear()
    get_document_service()
    interrupted = client.get(f"/api/projects/{project_id}/documents/{document_id}").json()
    assert interrupted["error_code"] == "INTERRUPTED"
    retried = client.post(f"/api/projects/{project_id}/documents/{document_id}/retry")
    assert retried.status_code == 200
    assert client.get(f"/api/projects/{project_id}/documents/{document_id}").json()["status"] == "ready"
    content_path = get_settings().data_dir / "documents" / project_id / document_id / "content.json"
    content_path.unlink()
    assert client.get(f"/api/projects/{project_id}/documents/{document_id}").json()["error_code"] == "MISSING_ARTIFACT"


def test_migration_keeps_existing_projects(tmp_path: Path):
    database = tmp_path / "old.sqlite3"
    with sqlite3.connect(database) as connection:
        connection.execute("CREATE TABLE projects (id TEXT PRIMARY KEY, name TEXT NOT NULL, description TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL)")
        connection.execute("INSERT INTO projects VALUES (?, ?, ?, ?, ?)", ("00000000-0000-0000-0000-000000000001", "Existing", None, "2026-01-01", "2026-01-01"))
    migrate(database)
    with sqlite3.connect(database) as connection:
        assert connection.execute("SELECT name FROM projects").fetchone()[0] == "Existing"
        assert connection.execute("SELECT name FROM sqlite_master WHERE name='documents'").fetchone()[0] == "documents"


def test_project_deletion_removes_only_its_document_files(client):
    first = project(client, "To delete")
    second = project(client, "Keep")
    removed = upload(client, first, "first.txt", b"Remove this document")["document"]
    kept = upload(client, second, "second.txt", b"Keep this document")["document"]
    removed_path = get_settings().data_dir / "documents" / first / removed["id"] / "original.txt"
    kept_path = get_settings().data_dir / "documents" / second / kept["id"] / "original.txt"
    assert client.delete(f"/api/projects/{first}").status_code == 204
    assert not removed_path.exists()
    assert kept_path.is_file()
    assert client.get(f"/api/projects/{second}/documents/{kept['id']}/content").status_code == 200
